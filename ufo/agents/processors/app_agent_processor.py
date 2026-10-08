# Copyright (c) Microsoft Corporation.
# Licensed under the MIT License.

import json
import os
from dataclasses import asdict, dataclass, field
from typing import TYPE_CHECKING, Any, ClassVar, Dict, List, Optional, Union

from pywinauto.controls.uiawrapper import UIAWrapper

import ufo.automator.ui_control.compound_operations.delete_duplicates
from ufo import utils
from ufo.agents.processors.actions import ActionSequence, BaseControlLog, OneStepAction
from ufo.agents.processors.basic import BaseProcessor
from ufo.automator.ui_control import ui_tree
from ufo.automator.ui_control.compound_operations.edit_xlsx_cells import EditXlsxCells
from ufo.automator.ui_control.compound_operations.select_xlsx_cell import SelectXlsxCell
from ufo.automator.ui_control.compound_operations.replace_previous_code import ReplacePreviousCode
from ufo.automator.ui_control.control_filter import ControlFilterFactory
from ufo.automator.ui_control.grounding.basic import BasicGrounding
from ufo.config.config import Config
from ufo.module.context import Context, ContextNames
from ufo.automator.ui_control.compound_operations.find_and_replace import FindAndReplace
from ufo.automator.ui_control.compound_operations.delete_duplicates import DeleteDuplicates
from ufo.automator.ui_control.compound_operations.check_available import CheckAvailable
from ufo.automator.ui_control.compound_operations.set_WPA_trace_filter import SetWPATraceFilter
from ufo.automator.ui_control.compound_operations.change_xxx_in_xxx_to_script import ChangeXXXinXXXtoScript
from ufo.automator.ui_control.compound_operations.move_cursor_to_beginning_of_page_xx import MoveCursorToBeginningOfPageXX
from ufo.automator.ui_control.compound_operations.paste_whole_column_to_txt import PasteWholeColumnToTxt

from ufo.prompts.graph import knowledge_graph_builder as kgb
from ufo.utils import print_with_color
from docx import Document
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pptx import Presentation
import base64
import re
from PIL import Image
from io import BytesIO

import win32com.client
from pathlib import Path
import psutil

if TYPE_CHECKING:
    from ufo.agents.agent.app_agent import AppAgent

configs = Config.get_instance().config_data

if configs is not None:
    CONTROL_BACKEND = configs.get("CONTROL_BACKEND", ["uia"])
    BACKEND = "win32" if "win32" in CONTROL_BACKEND else "uia"


@dataclass
class AppAgentAdditionalMemory:
    """
    The additional memory data for the AppAgent.
    """

    Step: int
    RoundStep: int
    AgentStep: int
    Round: int
    Subtask: str
    SubtaskIndex: int
    FunctionCall: List[str]
    Action: List[Dict[str, Any]]
    ActionSuccess: List[Dict[str, Any]]
    ActionType: List[str]
    Request: str
    Agent: str
    AgentName: str
    Application: str
    Cost: float
    Results: str
    error: str
    time_cost: Dict[str, float]
    ControlLog: Dict[str, Any]
    UserConfirm: Optional[str] = None


@dataclass
class AppAgentControlLog(BaseControlLog):
    """
    The control log data for the AppAgent.
    """

    control_friendly_class_name: str = ""
    control_coordinates: Dict[str, int] = field(default_factory=dict)


@dataclass
class ControlInfoRecorder:
    """
    The control meta information recorder for the current application window.
    """

    recording_fields: ClassVar[List[str]] = [
        "control_text",
        "control_type" if BACKEND == "uia" else "control_class",
        "control_rect",
        "source",
    ]

    application_windows_info: Dict[str, Any] = field(default_factory=dict)
    uia_controls_info: List[Dict[str, Any]] = field(default_factory=dict)
    grounding_controls_info: List[Dict[str, Any]] = field(default_factory=dict)
    merged_controls_info: List[Dict[str, Any]] = field(default_factory=dict)


@dataclass
class AppAgentRequestLog:
    """
    The request log data for the AppAgent.
    """

    step: int
    dynamic_examples: List[str]
    experience_examples: List[str]
    demonstration_examples: List[str]
    offline_docs: str
    online_docs: str
    dynamic_knowledge: str
    image_list: List[str]
    prev_subtask: List[str]
    plan: List[str]
    request: str
    control_info: List[Dict[str, str]]
    subtask: str
    current_application: str
    host_message: str
    blackboard_prompt: List[str]
    last_success_actions: List[Dict[str, Any]]
    include_last_screenshot: bool
    prompt: Dict[str, Any]
    control_info_recording: Dict[str, Any]


class AppAgentProcessor(BaseProcessor):
    """
    The processor for the app agent at a single step.
    """

    def __init__(
        self,
        agent: "AppAgent",
        context: Context,
        ground_service=Optional[BasicGrounding],
    ) -> None:
        """
        Initialize the app agent processor.
        :param agent: The app agent who executes the processor.
        :param context: The context of the session.
        """

        super().__init__(agent=agent, context=context)

        self.app_agent = agent
        self.host_agent = agent.host

        self._annotation_dict = None
        self._control_info = None
        self._operation = ""
        self._args = {}
        self._image_url = []
        self.control_filter_factory = ControlFilterFactory()
        self.control_recorder = ControlInfoRecorder()
        self.filtered_annotation_dict = None
        self.screenshot_save_path = None
        self.grounding_service = ground_service
        self.compound_operation_list = {"set_find_what_and_its_format", "set_replace_with_and_its_format"}
        self.knowledge_graph = kgb.KnowledgeGraph()
        # self.file_path,self.file_type = self.get_active_office_file()
        self.all_open_files = self.get_all_open_office_files()
        self.open_libreoffice_files = self.get_libreoffice_open_files()

    # @staticmethod
    # def get_active_office_file():
    #     """
    #     依次检测Word、Excel、PPT的活动文档路径
    #     返回第一个找到的有效文件路径
    #     """
    #     # 1. 先尝试检测Word
    #     try:
    #         word_app = win32com.client.GetObject(Class="Word.Application")
    #         if word_app and word_app.ActiveDocument:
    #             path = word_app.ActiveDocument.FullName
    #             if os.path.exists(path):
    #                 return path, "Word"
    #     except Exception:
    #         pass
    #
    #     # 2. 尝试检测Excel
    #     try:
    #         excel_app = win32com.client.GetObject(Class="Excel.Application")
    #         if excel_app and excel_app.ActiveWorkbook:
    #             path = excel_app.ActiveWorkbook.FullName
    #             if os.path.exists(path):
    #                 return path, "Excel"
    #     except Exception:
    #         pass
    #
    #     # 3. 尝试检测PowerPoint
    #     try:
    #         ppt_app = win32com.client.GetObject(Class="PowerPoint.Application")
    #         if ppt_app and ppt_app.ActivePresentation:
    #             path = ppt_app.ActivePresentation.FullName
    #             if os.path.exists(path):
    #                 return path, "PowerPoint"
    #     except Exception:
    #         pass
    #
    #     return None, None

    @staticmethod
    def get_all_open_office_files():
        """
        检测所有打开的Office文档（包括非活动文档）
        返回包含所有有效文件路径和应用类型的列表
        """
        import win32com.client
        import os
        import pythoncom

        results = []  # 存储所有找到的有效路径

        # 初始化COM环境（确保线程安全）
        pythoncom.CoInitialize()

        try:
            # 1. 检测所有Word文档
            try:
                word_app = win32com.client.GetObject(Class="Word.Application")
                if word_app:
                    # 遍历所有打开的文档（不只是活动文档）
                    for doc in word_app.Documents:
                        try:
                            path = doc.FullName
                            if os.path.exists(path):
                                results.append((path, "Word"))
                        except Exception:
                            continue
            except Exception:
                pass

            # 2. 检测所有Excel工作簿
            try:
                excel_app = win32com.client.GetObject(Class="Excel.Application")
                if excel_app:
                    # 遍历所有打开的工作簿（不只是活动工作簿）
                    for wb in excel_app.Workbooks:
                        try:
                            path = wb.FullName
                            if os.path.exists(path):
                                results.append((path, "Excel"))
                        except Exception:
                            continue
            except Exception:
                pass

            # 3. 检测所有PowerPoint演示文稿
            try:
                ppt_app = win32com.client.GetObject(Class="PowerPoint.Application")
                if ppt_app:
                    # 遍历所有打开的演示文稿（不只是活动演示文稿）
                    for pres in ppt_app.Presentations:
                        try:
                            path = pres.FullName
                            if os.path.exists(path):
                                results.append((path, "PowerPoint"))
                        except Exception:
                            continue
            except Exception:
                pass

            return results

        finally:
            # 清理COM资源
            pythoncom.CoUninitialize()

    @staticmethod
    def get_libreoffice_open_files() -> list[str]:
        WRITER_PROCESS = "soffice.exe"  # Writer 和 Calc 共用同一进程名
        CALC_PROCESS = "soffice.exe"

        # 目标扩展名
        TARGET_EXTENSIONS = {".odt", ".docx", ".xlsx"}
        """
        遍历所有 soffice.exe 进程，从命令行参数或打开的文件句柄中
        提取目标文档路径，返回去重后的路径列表。
        """
        found = []

        for proc in psutil.process_iter(["name", "cmdline", "open_files"]):
            try:
                name = (proc.info["name"] or "").lower()
                if name != "soffice.exe":
                    continue

                # ── 方法 A：从命令行参数解析路径 ──
                cmdline = proc.info["cmdline"] or []
                for arg in cmdline:
                    p = Path(arg)
                    if p.suffix.lower() in TARGET_EXTENSIONS and p.exists():
                        found.append(str(p.resolve()))

                # ── 方法 B：从进程打开的文件句柄中查找 ──
                # （需要管理员权限，否则 open_files 可能返回空）
                try:
                    open_files = proc.open_files()
                    for f in open_files:
                        p = Path(f.path)
                        if p.suffix.lower() in TARGET_EXTENSIONS and p.exists():
                            found.append(str(p.resolve()))
                except (psutil.AccessDenied, psutil.NoSuchProcess):
                    pass  # 无权限时跳过句柄检测，依赖方法 A

            except (psutil.NoSuchProcess, psutil.ZombieProcess, psutil.AccessDenied):
                continue

        # 去重，保持顺序
        seen = set()
        result = []
        for fp in found:
            if fp not in seen:
                seen.add(fp)
                result.append(fp)
        return result

    def print_step_info(self) -> None:
        """
        Print the step information.
        """
        utils.print_with_color(
            "Round {round_num}, Step {step}, AppAgent: Completing the subtask [{subtask}] on application [{application}].".format(
                round_num=self.round_num + 1,
                step=self.round_step + 1,
                subtask=self.subtask,
                application=self.application_process_name,
            ),
            "magenta",
        )

    def get_control_list(self, screenshot_path: str) -> List[UIAWrapper]:
        """
        Get the control list from the annotation dictionary.
        :param screenshot_path: The path to the clean screenshot.
        :return: The list of control items.
        """

        api_backend = None
        grounding_backend = None

        control_detection_backend = configs.get("CONTROL_BACKEND", ["uia"])

        if "uia" in control_detection_backend:
            api_backend = "uia"
        elif "win32" in control_detection_backend:
            api_backend = "win32"

        if "omniparser" in control_detection_backend:
            grounding_backend = "omniparser"

        if api_backend is not None:
            api_control_list = (
                self.control_inspector.find_control_elements_in_descendants(
                    self.application_window,
                    control_type_list=configs.get("CONTROL_LIST", []),
                    class_name_list=configs.get("CONTROL_LIST", []),
                )
            )
        else:
            api_control_list = []

        api_control_dict = {
            i + 1: control for i, control in enumerate(api_control_list)
        }

        # print(control_detection_backend, grounding_backend, screenshot_path)

        if grounding_backend == "omniparser" and self.grounding_service is not None:
            self.grounding_service: BasicGrounding

            onmiparser_configs = configs.get("OMNIPARSER", {})

            # print(onmiparser_configs)

            grounding_control_list = (
                self.grounding_service.convert_to_virtual_uia_elements(
                    image_path=screenshot_path,
                    application_window=self.application_window,
                    box_threshold=onmiparser_configs.get("BOX_THRESHOLD", 0.05),
                    iou_threshold=onmiparser_configs.get("IOU_THRESHOLD", 0.1),
                    use_paddleocr=onmiparser_configs.get("USE_PADDLEOCR", True),
                    imgsz=onmiparser_configs.get("IMGSZ", 640),
                )
            )
        else:
            grounding_control_list = []

        grounding_control_dict = {
            i + 1: control for i, control in enumerate(grounding_control_list)
        }

        merged_control_list = self.photographer.merge_control_list(
            api_control_list,
            grounding_control_list,
            iou_overlap_threshold=configs.get("IOU_THRESHOLD_FOR_MERGE", 0.1),
        )

        merged_control_dict = {
            i + 1: control for i, control in enumerate(merged_control_list)
        }

        # Record the control information for the uia controls.
        self.control_recorder.uia_controls_info = (
            self.control_inspector.get_control_info_list_of_dict(
                api_control_dict, ControlInfoRecorder.recording_fields
            )
        )

        # Record the control information for the grounding controls.
        self.control_recorder.grounding_controls_info = (
            self.control_inspector.get_control_info_list_of_dict(
                grounding_control_dict, ControlInfoRecorder.recording_fields
            )
        )

        # Record the control information for the merged controls.
        self.control_recorder.merged_controls_info = (
            self.control_inspector.get_control_info_list_of_dict(
                merged_control_dict, ControlInfoRecorder.recording_fields
            )
        )

        return merged_control_list

    @BaseProcessor.exception_capture
    @BaseProcessor.method_timer
    def capture_screenshot(self) -> None:
        """
        Capture the screenshot.
        """

        # Define the paths for the screenshots saved.
        screenshot_save_path = self.log_path + f"action_step{self.session_step}.png"
        self.screenshot_save_path = screenshot_save_path

        annotated_screenshot_save_path = (
            self.log_path + f"action_step{self.session_step}_annotated.png"
        )
        concat_screenshot_save_path = (
            self.log_path + f"action_step{self.session_step}_concat.png"
        )

        self._memory_data.add_values_from_dict(
            {
                "CleanScreenshot": screenshot_save_path,
                "AnnotatedScreenshot": annotated_screenshot_save_path,
                "ConcatScreenshot": concat_screenshot_save_path,
            }
        )

        self.photographer.capture_app_window_screenshot(
            self.application_window, save_path=screenshot_save_path
        )

        # Record the control information for the current application window.
        self.control_recorder.application_windows_info = (
            self.control_inspector.get_control_info(
                self.application_window, field_list=ControlInfoRecorder.recording_fields
            )
        )

        # # Get the control elements in the application window if the control items are not provided for reannotation.
        # if type(self.control_reannotate) == list and len(self.control_reannotate) > 0:
        #     control_list = self.control_reannotate
        # else:
        #     control_list = self.control_inspector.find_control_elements_in_descendants(
        #         self.application_window,
        #         control_type_list=configs.get("CONTROL_LIST", []),
        #         class_name_list=configs.get("CONTROL_LIST", []),
        #     )

        control_list = self.get_control_list(screenshot_save_path)

        # Get the annotation dictionary for the control items, in a format of {control_label: control_element}.
        self._annotation_dict = self.photographer.get_annotation_dict(
            self.application_window, control_list, annotation_type="number"
        )

        # Attempt to filter out irrelevant control items based on the previous plan.
        self.filtered_annotation_dict = self.get_filtered_annotation_dict(
            self._annotation_dict
        )

        # Capture the screenshot of the selected control items with annotation and save it.
        self.photographer.capture_app_window_screenshot_with_annotation_dict(
            self.application_window,
            self.filtered_annotation_dict,
            annotation_type="number",
            save_path=annotated_screenshot_save_path,
        )

        if configs.get("SAVE_UI_TREE", False):
            if self.application_window is not None:
                step_ui_tree = ui_tree.UITree(self.application_window)
                step_ui_tree.save_ui_tree_to_json(
                    os.path.join(
                        self.ui_tree_path, f"ui_tree_step{self.session_step}.json"
                    )
                )

        if configs.get("SAVE_FULL_SCREEN", False):

            desktop_save_path = (
                self.log_path + f"desktop_action_step{self.session_step}.png"
            )

            self._memory_data.add_values_from_dict(
                {"DesktopCleanScreenshot": desktop_save_path}
            )

            # Capture the desktop screenshot for all screens.
            self.photographer.capture_desktop_screen_screenshot(
                all_screens=True, save_path=desktop_save_path
            )

        # If the configuration is set to include the last screenshot with selected controls tagged, save the last screenshot.
        if configs.get("INCLUDE_LAST_SCREENSHOT", True):
            last_screenshot_save_path = (
                self.log_path + f"action_step{self.session_step - 1}.png"
            )
            last_control_screenshot_save_path = (
                self.log_path
                + f"action_step{self.session_step - 1}_selected_controls.png"
            )
            self._image_url += [
                self.photographer.encode_image_from_path(
                    last_control_screenshot_save_path
                    if os.path.exists(last_control_screenshot_save_path)
                    else last_screenshot_save_path
                )
            ]

        # Whether to concatenate the screenshots of clean screenshot and annotated screenshot into one image.
        if configs.get("CONCAT_SCREENSHOT", False):
            self.photographer.concat_screenshots(
                screenshot_save_path,
                annotated_screenshot_save_path,
                concat_screenshot_save_path,
            )
            self._image_url += [
                self.photographer.encode_image_from_path(concat_screenshot_save_path)
            ]
        else:
            screenshot_url = self.photographer.encode_image_from_path(
                screenshot_save_path
            )
            screenshot_annotated_url = self.photographer.encode_image_from_path(
                annotated_screenshot_save_path
            )
            self._image_url += [screenshot_url, screenshot_annotated_url]

        # Save the XML file for the current state.
        if configs.get("LOG_XML", False):
            self._save_to_xml()

    @BaseProcessor.exception_capture
    @BaseProcessor.method_timer
    def get_control_info(self) -> None:
        """
        Get the control information.
        """

        # Get the control information for the control items and the filtered control items, in a format of list of dictionaries.
        self._control_info = self.control_inspector.get_control_info_list_of_dict(
            self._annotation_dict,
            ["control_text", "control_type" if BACKEND == "uia" else "control_class"],
        )
        self.filtered_control_info = (
            self.control_inspector.get_control_info_list_of_dict(
                self.filtered_annotation_dict,
                [
                    "control_text",
                    "control_type" if BACKEND == "uia" else "control_class",
                ],
            )
        )

    @BaseProcessor.exception_capture
    @BaseProcessor.method_timer
    def get_prompt_message(self) -> None:
        """
        Get the prompt message for the AppAgent.
        """

        experience_results, demonstration_results = (
            self.app_agent.demonstration_prompt_helper(request=self.subtask)
        )

        retrieved_results = experience_results + demonstration_results

        # Get the external knowledge prompt for the AppAgent using the offline and online retrievers.

        offline_docs, online_docs = self.app_agent.external_knowledge_prompt_helper(
            self.subtask,
            configs.get("RAG_OFFLINE_DOCS_RETRIEVED_TOPK", 0),
            configs.get("RAG_ONLINE_RETRIEVED_TOPK", 0),
        )

        # print(offline_docs, online_docs)

        external_knowledge_prompt = offline_docs + online_docs

        if not self.app_agent.blackboard.is_empty():
            blackboard_prompt = self.app_agent.blackboard.blackboard_to_prompt()
        else:
            blackboard_prompt = []

        # Get the last successful actions of the AppAgent.
        last_success_actions = self.get_last_success_actions()

        action_keys = ["Function", "Args", "ControlText", "Results", "RepeatTimes"]

        filtered_last_success_actions = [
            {key: action.get(key, "") for key in action_keys}
            for action in last_success_actions
        ]

        def read_word_file(file_path):
            doc = Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            # 如果有表格，也可以读取
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)
            return '\n'.join(full_text)
        #
        # file_content = read_word_file(self.file_path)

        # def read_excel_to_prompt(file_path):
        #     # """
        #     # 读取XLSX文件内容并格式化为提示词字符串
        #     workbook = load_workbook(file_path, read_only=True)
        #     sheet = workbook.active
        #
        #     excel_content = []
        #     for row_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
        #         for col_idx, value in enumerate(row, 1):
        #             if value is not None:  # 跳过空单元格
        #                 # coord = f"{get_column_letter(col_idx)}{row_idx}"
        #                 coord = f"{'column ' + get_column_letter(col_idx)} row {row_idx}"
        #                 excel_content.append(f"{coord}: {value}")
        #
        #     workbook.close()  # 显式关闭工作簿
        #     return "\n".join(excel_content)
        def read_excel_to_prompt(file_path):
            workbook = load_workbook(file_path, read_only=True)
            sheet = workbook.active

            rows_data = []
            start_row = None
            start_col = None

            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
                for col_idx, cell in enumerate(row, 1):
                    if cell is not None:
                        if start_row is None:
                            start_row = row_idx
                            start_col = col_idx
                rows_data.append([cell if cell is not None else "" for cell in row])

            workbook.close()

            if not rows_data or start_row is None:
                return ""

            # 裁剪掉起始位置之前的行和列
            rows_data = [row[start_col - 1:] for row in rows_data[start_row - 1:]]

            col_count = len(rows_data[0])
            row_count = len(rows_data)

            # 生成列标题行：第一格留空，后面是 A/B/C... 列号
            col_letters = [get_column_letter(start_col + i) for i in range(col_count)]
            header_row = "| " + " | ".join([""] + col_letters) + " |"
            separator = "| " + " | ".join(["---"] * (col_count + 1)) + " |"

            # 每行前面加行号
            lines = [header_row, separator]
            for i, row in enumerate(rows_data):
                row_num = start_row + i
                lines.append("| " + " | ".join([str(row_num)] + [str(c) for c in row]) + " |")

            return "\n".join(lines)


        def extract_ppt_content(file_path):
            """
            从PPT文件中提取文本、图片和结构信息
            返回适合大模型输入的格式化内容
            """
            prs = Presentation(file_path)
            content = {"slides": [], "metadata": {}}

            # 提取元数据
            content["metadata"] = {
                "title": prs.core_properties.title,
                "author": prs.core_properties.author,
                "created": str(prs.core_properties.created),
                "modified": str(prs.core_properties.modified),
                "slide_count": len(prs.slides)
            }

            # 提取每页幻灯片内容
            for i, slide in enumerate(prs.slides):
                slide_content = {
                    "slide_number": i + 1,
                    "title": "",
                    "text": [],
                    "images": [],
                    "notes": ""
                }

                # 提取标题
                if slide.shapes.title:
                    slide_content["title"] = slide.shapes.title.text

                # 提取文本框内容
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape != slide.shapes.title:
                        text = shape.text.strip()
                        if text:
                            slide_content["text"].append(text)

                # 提取图片（转换为Base64）
                for shape in slide.shapes:
                    if shape.shape_type == 13:  # 13 表示图片类型
                        # img_data = BytesIO()
                        # shape.image.blob.seek(0)
                        # img_data.write(shape.image.blob.read())
                        # img = Image.open(img_data)
                        #
                        # # 转换为Base64
                        # buffered = BytesIO()
                        # img.save(buffered, format="PNG")
                        # img_base64 = base64.b64encode(buffered.getvalue()).decode("utf-8")
                        img_data = shape.image.blob
                        img_base64 = base64.b64encode(img_data).decode('utf-8')

                        slide_content["images"].append({
                            "description": f"Slide {i + 1} Image {len(slide_content['images']) + 1}",
                            "base64": img_base64
                        })

                # 提取演讲者备注
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    slide_content["notes"] = slide.notes_slide.notes_text_frame.text.strip()

                content["slides"].append(slide_content)

            return content

        def format_for_prompt(content, include_images=True):
            """
            将提取的内容格式化为适合大模型的提示词
            """
            prompt = f"PPT文档分析报告\n"
            prompt += f"标题: {content['metadata']['title']}\n"
            prompt += f"作者: {content['metadata']['author']}\n"
            prompt += f"创建日期: {content['metadata']['created']}\n"
            prompt += f"最后修改: {content['metadata']['modified']}\n"
            prompt += f"总页数: {content['metadata']['slide_count']}\n\n"

            # 添加每页内容
            for slide in content["slides"]:
                prompt += f"### 第 {slide['slide_number']} 页: {slide['title']}\n"

                # 添加文本内容
                if slide["text"]:
                    prompt += "文本内容:\n"
                    for i, text in enumerate(slide["text"], 1):
                        prompt += f"  {i}. {clean_text(text)}\n"

                # 添加图片（如果需要）
                if include_images and slide["images"]:
                    prompt += "图片描述:\n"
                    for img in slide["images"]:
                        prompt += f"  - {img['description']}\n"
                        # 如果大模型支持Base64图像输入，可以直接包含
                        # prompt += f"  - [图像数据: {img['base64'][:30]}...]\n"

                # 添加备注
                if slide["notes"]:
                    prompt += f"演讲者备注: {clean_text(slide['notes'])}\n"

                prompt += "\n"

            return prompt

        def clean_text(text):
            """清理文本中的特殊字符和多余空格"""
            # 移除PPT中的特殊控制字符
            text = re.sub(r'[\x00-\x1F\x7F-\x9F]', ' ', text)
            # 合并多个空格
            text = re.sub(r'\s+', ' ', text).strip()
            return text

        # # 使用示例
        # if __name__ == "__main__":
        #     ppt_path = "your_presentation.pptx"  # 替换为你的PPT文件路径
        #
        #     # 1. 提取PPT内容
        #     ppt_content = extract_ppt_content(ppt_path)
        #
        #     # 2. 格式化为提示词
        #     prompt = format_for_prompt(ppt_content, include_images=True)
        #
        #     print("=" * 50)
        #     print("生成的提示词内容:")
        #     print("=" * 50)
        #     print(prompt)

        successors = list(self.knowledge_graph.graph.successors(self.context.current_node))
        file_content = []
        for i in range(len(successors)):
            if "Read_task_related_file" in successors[i]:
                for i in range(len(self.all_open_files)):
                    if self.all_open_files[i][1] == 'Word':
                        file_content.append({
                            "type": "document",
                            "source": {
                                "type": "text",
                                "media_type": "docx",
                                "file_name": self.all_open_files[i][0][self.all_open_files[i][0].rfind('\\') + 1:
                                                                       self.all_open_files[i][
                                                                           0].rfind('.')],
                                "data": read_word_file(self.all_open_files[i][0]) + '\n'
                                    },
                                             })
                    elif self.all_open_files[i][1] == 'Excel':
                        file_content.append({
                            "type": "document",
                            "source": {
                                "type": "text",
                                "media_type": "xlsx",
                                "file_name": self.all_open_files[i][0][self.all_open_files[i][0].rfind('\\') + 1:
                                                                       self.all_open_files[i][
                                                                           0].rfind('.')],
                                "data": "以下是Excel文件中所包含的有效单元格的control_text（例如B6）及其对应内容(例如6/4/2022)，根据任务需求选择目标单元格对应的ControlText（例如B6）进行操作" + ':\n' + read_excel_to_prompt(self.all_open_files[i][0]) + '\n'
                            },
                        })
                    elif self.all_open_files[i][1] == 'PowerPoint':
                        file_content_unformatted = extract_ppt_content(self.all_open_files[i][0])
                        file_content.append({
                            "type": "document",
                            "source": {
                                "type": "text",
                                "media_type": "ppt",
                                "file_name": self.all_open_files[i][0][self.all_open_files[i][0].rfind('\\') + 1:
                                                                       self.all_open_files[i][
                                                                           0].rfind('.')],
                                "data": format_for_prompt(file_content_unformatted, include_images=True) + '\n'
                            },
                        })

                for i in range(len(self.open_libreoffice_files)):
                    if 'docx' in self.open_libreoffice_files[i]:
                        file_content.append({
                            "type": "document",
                            "source": {
                                "type": "text",
                                "media_type": "docx",
                                "file_name": self.open_libreoffice_files[i][self.open_libreoffice_files[i].rfind('\\') + 1:
                                                                       self.open_libreoffice_files[i].rfind('.')],
                                "data": "Below are file/files content directly related to the task, use it to decide the next action. " + ':\n' + read_word_file(self.open_libreoffice_files[i]) + '\n'
                            },
                        })

                    if 'xlsx' in self.open_libreoffice_files[i]:
                        file_content.append({
                            "type": "document",
                            "source": {
                                "type": "text",
                                "media_type": "xlsx",
                                "file_name": self.open_libreoffice_files[i][self.open_libreoffice_files[i].rfind('\\') + 1:
                                                                            self.open_libreoffice_files[i].rfind('.')],
                                "data": "以下是xlsx文件中所包含的有效单元格，单元格对应的列号和行号显示在Markdown顶部和左侧，根据任务需求选择目标单元格（例如B6）进行操作" + ':\n' + read_excel_to_prompt(self.open_libreoffice_files[i]) + '\n'
                            },
                        })
                break


                # file_content += "Following document (PowerPoint) related to task with file name " + self.all_open_files[i][0][
                #                                                                               self.all_open_files[i][
                #                                                                                   0].rfind('\\') + 1:
                #                                                                               self.all_open_files[i][
                #                                                                                   0].rfind('.')] + ':\n'
                # file_content_unformatted = extract_ppt_content(self.all_open_files[i][0])
                # file_content += format_for_prompt(file_content_unformatted, include_images=True)
                # file_content += '\n'
            # else:
            #     file_content += ''

        # def load_files_from_folder(folder_path: str) -> list[dict]:
        #     """
        #     扫描文件夹，读取所有 JPEG/JPG/PDF 文件并编码为 base64。
        #
        #     返回值格式（每个文件一个 dict）：
        #         {
        #             "filename": "example.jpg",
        #             "media_type": "image/jpeg",  # 或 "application/pdf"
        #             "data": "<base64字符串>",
        #         }
        #     """
        #     SUPPORTED_EXTENSIONS = {
        #         ".jpg": "image/jpeg",
        #         ".jpeg": "image/jpeg",
        #         ".pdf": "application/pdf",
        #     }
        #     folder = Path(folder_path)
        #     if not folder.is_dir():
        #         raise NotADirectoryError(f"路径不是有效文件夹：{folder_path}")
        #
        #     files_content = []
        #
        #     for file_path in sorted(folder.iterdir()):
        #         ext = file_path.suffix.lower()
        #         if ext not in SUPPORTED_EXTENSIONS:
        #             continue  # 跳过不支持的格式
        #
        #         media_type = SUPPORTED_EXTENSIONS[ext]
        #         with open(file_path, "rb") as f:
        #             encoded = base64.standard_b64encode(f.read()).decode("utf-8")
        #
        #         files_content.append({
        #             "filename": file_path.name,
        #             "media_type": media_type,
        #             "data": encoded,
        #         })
        #         print(f"  ✓ 已加载：{file_path.name}  ({media_type})")
        #
        #     return files_content
        #
        # image_file_content = load_files_from_folder("D:\QD\pycharmproject\OSWorld-main\OSWorld-main\evaluation_examples\examples_windows\multi_app\multi_app\8e116af7-7db7-4e35-a68b-b0939c066c78")
        # # file_content += "Following are images related to task" + ':\n' + image_file_content
        # file_content += image_file_content

        # Construct the prompt message for the AppAgent.
        self._prompt_message = self.app_agent.message_constructor(
            dynamic_examples=retrieved_results,
            dynamic_knowledge=external_knowledge_prompt,
            image_list=self._image_url,
            control_info=self.filtered_control_info,
            prev_subtask=self.previous_subtasks,
            plan=self.prev_plan,
            request=self.request,
            subtask=self.subtask,
            current_application=self.application_process_name,
            host_message=self.host_message,
            blackboard_prompt=blackboard_prompt,
            last_success_actions=filtered_last_success_actions,
            include_last_screenshot=configs.get("INCLUDE_LAST_SCREENSHOT", True),
            current_node=self.context.current_node,
            file_content=file_content
        )

        # Log the prompt message. Only save them in debug mode.
        request_data = AppAgentRequestLog(
            step=self.session_step,
            experience_examples=experience_results,
            demonstration_examples=demonstration_results,
            dynamic_examples=retrieved_results,
            offline_docs=offline_docs,
            online_docs=online_docs,
            dynamic_knowledge=external_knowledge_prompt,
            image_list=self._image_url,
            prev_subtask=self.previous_subtasks,
            plan=self.prev_plan,
            request=self.request,
            control_info=self.filtered_control_info,
            subtask=self.subtask,
            current_application=self.application_process_name,
            host_message=self.host_message,
            blackboard_prompt=blackboard_prompt,
            last_success_actions=filtered_last_success_actions,
            include_last_screenshot=configs.get("INCLUDE_LAST_SCREENSHOT", True),
            prompt=self._prompt_message,
            control_info_recording=asdict(self.control_recorder),
        )

        request_log_str = json.dumps(asdict(request_data), ensure_ascii=False)
        self.request_logger.debug(request_log_str)

    @BaseProcessor.exception_capture
    @BaseProcessor.method_timer
    def get_response(self) -> None:
        """
        Get the response from the LLM.
        """
        successors = list(self.knowledge_graph.graph.successors(self._context.current_node))
        if len(successors) == 1 and "xx" not in successors[0]:
            self._response = '{\n  "Observation": "",\n  "Thought": "",\n  "ControlLabel": "",\n  "ControlText": "",\n  "SelectedAction": 1,\n  "Args": "",\n  "Status": "CONTINUE",\n  "Plan": [],\n  "Comment": "",\n  "SaveScreenshot": {"save": false}\n}'
            self.cost = 0
        else:
            retry = 0
            while retry < configs.get("JSON_PARSING_RETRY", 3):
                # Try to get the response from the LLM. If an error occurs, catch the exception and log the error.
                self._response, self.cost = self.app_agent.get_response(
                    self._prompt_message, "APPAGENT", use_backup_engine=True
                )

                try:
                    self.app_agent.response_to_dict(self._response)
                    break
                except Exception as e:
                    print("Error in parsing response: ", e)
                    retry += 1

    @BaseProcessor.exception_capture
    @BaseProcessor.method_timer
    def parse_response(self) -> None:
        """
        Parse the response.
        """

        self._response_json = self.app_agent.response_to_dict(self._response)

        # self.control_label = self._response_json.get("ControlLabel", "")
        # self.control_text = self._response_json.get("ControlText", "")
        successors = list(self.knowledge_graph.graph.successors(self._context.current_node))
        if successors and self._response_json.get('SelectedAction'):
            self._context.current_node = successors[int(self._response_json.get('SelectedAction')) - 1]
        if not list(self.knowledge_graph.graph.successors(self._context.current_node)):
            self._response_json['Status'] = 'FINISH'
        self.status = self._response_json.get("Status", "")
        selected_action = eval(self._context.current_node)
        if 'ControlText' in selected_action and selected_action['ControlText'] == "xx":
            selected_action['ControlText'] = self._response_json['ControlText']

        if 'ControlText' in selected_action and 'ControlLabel' not in selected_action:
            for i, item in enumerate(self.filtered_control_info):
                if item.get('control_text') == selected_action['ControlText']:
                    selected_action['ControlLabel'] = str(i + 1)
                    break
            else:
                selected_action['ControlLabel'] = ''
                print_with_color("Control Text not Found in List", "red")

        if '#th UI after control text ##' in selected_action and 'ControlLabel' not in selected_action:
            for i, item in enumerate(self.filtered_control_info):
                if item.get('control_text') == selected_action['control text ##']:
                    selected_action['ControlLabel'] = str(i + 1 + selected_action['#th UI after control text ##'])
                    selected_action['ControlText'] = self.filtered_annotation_dict[selected_action['ControlLabel']].element_info.name
                    break
            else:
                selected_action['ControlLabel'] = ''
                print_with_color("Control Text not Found in List", "red")

        if '#th UI after control text contain ##' in selected_action and 'ControlLabel' not in selected_action:
            for i, item in enumerate(self.filtered_control_info):
                if selected_action['control text contain ##'] in item.get('control_text'):
                    selected_action['ControlLabel'] = str(i + 1 + selected_action['#th UI after control text contain ##'])
                    selected_action['ControlText'] = self.filtered_annotation_dict[selected_action['ControlLabel']].element_info.name
                    break
            else:
                selected_action['ControlLabel'] = ''
                print_with_color("Control Text not Found in List", "red")

        if 'ControlText_Contain' in selected_action and 'ControlLabel' not in selected_action:
            for i, item in enumerate(self.filtered_control_info):
                if selected_action['ControlText_Contain'] in item.get('control_text'):
                    selected_action['ControlLabel'] = str(i + 1)
                    selected_action['ControlText'] = self.filtered_annotation_dict[str(i + 1)].element_info.name
                    break
            else:
                selected_action['ControlLabel'] = ''
                print_with_color("Control Text not Found in List", "red")

        if 'Previous_ControlText_Contain' in selected_action and 'ControlLabel' not in selected_action:
            for i, item in enumerate(self.filtered_control_info):
                if selected_action['Previous_ControlText_Contain'] in item.get('control_text'):
                    selected_action['ControlLabel'] = str(i + 2)
                    selected_action['ControlText'] = self.filtered_annotation_dict[str(i + 2)].element_info.name
                    break
            else:
                selected_action['ControlLabel'] = ''
                print_with_color("Control Text not Found in List", "red")

        if 'ControlText_contain_list' in selected_action:
            for i in range(len(self.filtered_annotation_dict)):
                if any(substring in self.filtered_annotation_dict[str(i + 1)].element_info.name for substring in
                       selected_action['ControlText_contain_list']):
                    self._response_json['ControlText'] = self.filtered_annotation_dict[str(i + 1)].element_info.name
                    self._response_json["ControlLabel"] = str(i + 1)
                    break
            else:
                # selected_action['ControlText'] = ''
                print_with_color("Control Text not Found in List", "red")

        # if 'xx' in self._context.current_node or 'xxx' in self._context.current_node:
        #     if 'text' in selected_action['Args']:
        #         if 'text' in self._response_json['Args']:
        #             selected_action['Args']['text'] = self._response_json['Args']['text']
        #         else:
        #             selected_action['Args']['text'] = self._response_json['Args']['keys']
        #     elif 'keys' in selected_action['Args']:
        #         if 'text' in self._response_json['Args']:
        #             selected_action['Args']['keys'] = self._response_json['Args']['text']
        #         else:
        #             selected_action['Args']['keys'] = self._response_json['Args']['keys']
        #     else:
        #         del selected_action['Args']
        if 'Args' in selected_action:
            for i, item in enumerate(selected_action['Args']):
                if selected_action['Args'][item] == 'xx' or selected_action['Args'][item] == 'xxx':
                    selected_action['Args'][item] = self._response_json['Args'][item]

        # if 'ControlText' in selected_action and selected_action['ControlText'] == "xx":
        #     del selected_action['ControlText']

        self._response_json.update(selected_action)

        self.control_text = self._response_json.get("ControlText", "")
        self.control_label = self._response_json.get("ControlLabel", "")

        self._operation = self._response_json.get("Function", "")
        self.question_list = self._response_json.get("Questions", [])
        self._args = utils.revise_line_breaks(self._response_json.get("Args", ""))

        # Convert the plan from a string to a list if the plan is a string.
        self.plan = self.string2list(self._response_json.get("Plan", ""))
        self._response_json["Plan"] = self.plan

        self.status = self._response_json.get("Status", "")
        self.app_agent.print_response(
            response_dict=self._response_json, print_action=True
        )

        if not list(self.knowledge_graph.graph.successors(self._context.current_node)):
            self._context.current_node = "begin"
        elif self.status == 'FINISH' and "Reference_task" not in list(self.knowledge_graph.graph.successors(self._context.current_node))[0]:
            self._context.current_node = "begin"


    @BaseProcessor.exception_capture
    @BaseProcessor.method_timer
    def execute_action(self) -> None:
        """
        Execute the action.
        """
        if self.application_window.element_info.name == "查找和替换" and self._operation in ["set_find_what_and_its_format", "set_replace_with_and_its_format"]:
            fr = FindAndReplace(self)
            if self._operation == "set_find_what_and_its_format":
                # 设置查找内容和格式
                fr.set_find_what_and_its_format()
            else:
                # 设置替换内容和格式
                fr.set_replace_with_and_its_format()
        elif self._operation == "delete_duplicates":
            dd=DeleteDuplicates(self)
            dd.delete_duplicates()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])
        elif self._operation == "change_xxx_in_xxx_to_subscript":
            cx=ChangeXXXinXXXtoScript(self)
            cx.change_xxx_in_xxx_to_subscript()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])
        elif self._operation == "move_cursor_to_beginning_of_page_xx":
            mc=MoveCursorToBeginningOfPageXX(self)
            mc.move_cursor_to_beginning_of_page_xx()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])

        elif self._operation == "check_available":
            ca=CheckAvailable(self)
            ca.check_available()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])
        elif self._operation == "set_WPA_trace_filter":
            stf=SetWPATraceFilter(self)
            stf.set_WPA_trace_filter()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])
        elif self._operation == "paste_whole_column_to_txt":
            pwc=PasteWholeColumnToTxt(self)
            pwc.paste_whole_column_to_txt()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])
        elif self._operation == "replace_previous_code":
            rpc=ReplacePreviousCode(self)
            rpc.replace_previous_code()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])
        elif self._operation == "edit_xlsx_cells":
            exc=EditXlsxCells(self)
            exc.edit_xlsx_cells()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])
        elif self._operation == "select_xlsx_cell":
            sxc=SelectXlsxCell(self)
            sxc.select_xlsx_cell()
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            self.actions: ActionSequence = ActionSequence(actions=[action])
        else:
            action = OneStepAction(
                function=self._operation,
                args=self._args,
                control_label=self._control_label,
                control_text=self.control_text,
                after_status=self.status,
            )
            control_selected = self._annotation_dict.get(self._control_label, None)

            # Save the screenshot of the tagged selected control.
            self.capture_control_screenshot(control_selected)

            self.actions: ActionSequence = ActionSequence(actions=[action])

            self.actions.execute_all(
                puppeteer=self.app_agent.Puppeteer,
                control_dict=self._annotation_dict,
                application_window=self.application_window,
            )

            if self.is_application_closed():
                utils.print_with_color("Warning: The application is closed.", "yellow")
                self.status = "FINISH"

    def capture_control_screenshot(
        self, control_selected: Union[UIAWrapper, List[UIAWrapper]]
    ) -> None:
        """
        Capture the screenshot of the selected control.
        :param control_selected: The selected control item or a list of selected control items.
        """
        control_screenshot_save_path = (
            self.log_path + f"action_step{self.session_step}_selected_controls.png"
        )

        self._memory_data.add_values_from_dict(
            {"SelectedControlScreenshot": control_screenshot_save_path}
        )

        sub_control_list = (
            control_selected
            if isinstance(control_selected, list)
            else [control_selected]
        )

        self.photographer.capture_app_window_screenshot_with_rectangle(
            self.application_window,
            sub_control_list=sub_control_list,
            save_path=control_screenshot_save_path,
            background_screenshot_path=self.screenshot_save_path,
        )

    def handle_screenshot_status(self) -> None:
        """
        Handle the screenshot status when the annotation is overlapped and the agent is unable to select the control items.
        """

        utils.print_with_color(
            "Annotation is overlapped and the agent is unable to select the control items. New annotated screenshot is taken.",
            "magenta",
        )
        self.control_reannotate = self.app_agent.Puppeteer.execute_command(
            "annotation", self._args, self._annotation_dict
        )

    def sync_memory(self):
        """
        Sync the memory of the AppAgent.
        """

        app_root = self.control_inspector.get_application_root_name(
            self.application_window
        )

        action_type = [
            self.app_agent.Puppeteer.get_command_types(action.function)
            for action in self.actions.actions
        ]

        all_previous_success_actions = self.get_all_success_actions()

        action_success = self.actions.to_list_of_dicts(
            success_only=True, previous_actions=all_previous_success_actions
        )

        # Create the additional memory data for the log.
        additional_memory = AppAgentAdditionalMemory(
            Step=self.session_step,
            RoundStep=self.round_step,
            AgentStep=self.app_agent.step,
            Round=self.round_num,
            Subtask=self.subtask,
            SubtaskIndex=self.round_subtask_amount,
            FunctionCall=self.actions.get_function_calls(),
            Action=self.actions.to_list_of_dicts(
                previous_actions=all_previous_success_actions
            ),
            ActionSuccess=action_success,
            ActionType=action_type,
            Request=self.request,
            Agent="AppAgent",
            AgentName=self.app_agent.name,
            Application=app_root,
            Cost=self._cost,
            Results=self.actions.get_results(),
            error=self._exeception_traceback,
            time_cost=self._time_cost,
            ControlLog=self.actions.get_control_logs(),
            UserConfirm=(
                "Yes"
                if self.status.upper()
                == self._agent_status_manager.CONFIRM.value.upper()
                else None
            ),
        )

        # Log the original response from the LLM.
        self.add_to_memory(self._response_json)

        # Log the additional memory data for the AppAgent.
        self.add_to_memory(asdict(additional_memory))

    def update_memory(self) -> None:
        """
        Update the memory of the Agent.
        """

        # Sync the memory of the AppAgent.
        self.sync_memory()

        self.app_agent.add_memory(self._memory_data)

        # Log the memory item.
        self.context.add_to_structural_logs(self._memory_data.to_dict())
        # self.log(self._memory_data.to_dict())

        # Only memorize the keys in the HISTORY_KEYS list to feed into the prompt message in the future steps.
        memorized_action = {
            key: self._memory_data.to_dict().get(key)
            for key in configs.get("HISTORY_KEYS", [])
        }

        if self.is_confirm():

            if self._is_resumed:
                self._memory_data.add_values_from_dict({"UserConfirm": "Yes"})
                memorized_action["UserConfirm"] = "Yes"
            else:
                self._memory_data.add_values_from_dict({"UserConfirm": "No"})
                memorized_action["UserConfirm"] = "No"

        # Save the screenshot to the blackboard if the SaveScreenshot flag is set to True by the AppAgent.
        self._update_image_blackboard()
        self.host_agent.blackboard.add_trajectories(memorized_action)

    def get_all_success_actions(self) -> List[Dict[str, Any]]:
        """
        Get the previous action.
        :return: The previous action of the agent.
        """
        agent_memory = self.app_agent.memory

        if agent_memory.length > 0:
            success_action_memory = agent_memory.filter_memory_from_keys(
                ["ActionSuccess"]
            )
            success_actions = []
            for success_action in success_action_memory:
                success_actions += success_action.get("ActionSuccess", [])

        else:
            success_actions = []

        return success_actions

    def get_last_success_actions(self) -> List[Dict[str, Any]]:
        """
        Get the previous action.
        :return: The previous action of the agent.
        """
        agent_memory = self.app_agent.memory

        if agent_memory.length > 0:
            last_success_actions = (
                agent_memory.get_latest_item().to_dict().get("ActionSuccess", [])
            )

        else:
            last_success_actions = []

        return last_success_actions

    def _update_image_blackboard(self) -> None:
        """
        Save the screenshot to the blackboard if the SaveScreenshot flag is set to True by the AppAgent.
        """
        screenshot_saving = self._response_json.get("SaveScreenshot", {})

        if screenshot_saving.get("save", False):

            screenshot_save_path = self.log_path + f"action_step{self.session_step}.png"
            metadata = {
                "screenshot application": self.context.get(
                    ContextNames.APPLICATION_PROCESS_NAME
                ),
                "saving reason": screenshot_saving.get("reason", ""),
            }
            self.app_agent.blackboard.add_image(screenshot_save_path, metadata)

    def _save_to_xml(self) -> None:
        """
        Save the XML file for the current state. Only work for COM objects.
        """
        log_abs_path = os.path.abspath(self.log_path)
        xml_save_path = os.path.join(
            log_abs_path, f"xml/action_step{self.session_step}.xml"
        )
        self.app_agent.Puppeteer.save_to_xml(xml_save_path)

    def get_filtered_annotation_dict(
        self, annotation_dict: Dict[str, UIAWrapper], configs: Dict[str, Any] = configs
    ) -> Dict[str, UIAWrapper]:
        """
        Get the filtered annotation dictionary.
        :param annotation_dict: The annotation dictionary.
        :return: The filtered annotation dictionary.
        """

        # Get the control filter type and top k plan from the configuration.
        control_filter_type = configs["CONTROL_FILTER_TYPE"]
        topk_plan = configs["CONTROL_FILTER_TOP_K_PLAN"]

        if len(control_filter_type) == 0 or self.prev_plan == []:
            return annotation_dict

        control_filter_type_lower = [
            control_filter_type_lower.lower()
            for control_filter_type_lower in control_filter_type
        ]

        filtered_annotation_dict = {}

        # Get the top k recent plans from the memory.
        plans = self.control_filter_factory.get_plans(self.prev_plan, topk_plan)

        # Filter the annotation dictionary based on the keywords of control text and plan.
        if "text" in control_filter_type_lower:
            model_text = self.control_filter_factory.create_control_filter("text")
            filtered_text_dict = model_text.control_filter(annotation_dict, plans)
            filtered_annotation_dict = (
                self.control_filter_factory.inplace_append_filtered_annotation_dict(
                    filtered_annotation_dict, filtered_text_dict
                )
            )

        # Filter the annotation dictionary based on the semantic similarity of the control text and plan with their embeddings.
        if "semantic" in control_filter_type_lower:
            model_semantic = self.control_filter_factory.create_control_filter(
                "semantic", configs["CONTROL_FILTER_MODEL_SEMANTIC_NAME"]
            )
            filtered_semantic_dict = model_semantic.control_filter(
                annotation_dict, plans, configs["CONTROL_FILTER_TOP_K_SEMANTIC"]
            )
            filtered_annotation_dict = (
                self.control_filter_factory.inplace_append_filtered_annotation_dict(
                    filtered_annotation_dict, filtered_semantic_dict
                )
            )

        # Filter the annotation dictionary based on the icon image icon and plan with their embeddings.
        if "icon" in control_filter_type_lower:
            model_icon = self.control_filter_factory.create_control_filter(
                "icon", configs["CONTROL_FILTER_MODEL_ICON_NAME"]
            )

            cropped_icons_dict = self.photographer.get_cropped_icons_dict(
                self.application_window, annotation_dict
            )
            filtered_icon_dict = model_icon.control_filter(
                annotation_dict,
                cropped_icons_dict,
                plans,
                configs["CONTROL_FILTER_TOP_K_ICON"],
            )
            filtered_annotation_dict = (
                self.control_filter_factory.inplace_append_filtered_annotation_dict(
                    filtered_annotation_dict, filtered_icon_dict
                )
            )

        return filtered_annotation_dict
